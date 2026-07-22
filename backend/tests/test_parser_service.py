"""Tests for the document parser service."""

import tempfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from app.services.parser_service import parse_docx, parse_pdf, parse_xlsx


def _create_test_pdf(text: str) -> str:
    """Create a temporary PDF with *text* on a single page and return its path."""
    import fitz
    tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), text, fontsize=12)
    doc.save(tmp.name)
    doc.close()
    return tmp.name


def _create_test_docx(paragraphs: list[str]) -> str:
    """Create a temporary DOCX and return its path."""
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(tmp.name)
    return tmp.name


def _create_test_xlsx(sheets: dict[str, list[list[str]]]) -> str:
    """Create a temporary XLSX from *sheets* mapping name → rows."""
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb = Workbook()
    first = True
    for sheet_name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(title=sheet_name)
        for row in rows:
            ws.append(row)
    wb.save(tmp.name)
    return tmp.name


class TestParsePdf:
    """RED→GREEN→REFACTOR: PDF text extraction."""

    @pytest.mark.asyncio
    async def test_extracts_text(self):
        """GIVEN a PDF with text WHEN parsed THEN text is extracted."""
        path = _create_test_pdf("Este es un texto de prueba PDF.")
        try:
            result = await parse_pdf(path)
            assert len(result) >= 1
            combined = "\n".join(r["text"] for r in result)
            assert "Este es un texto de prueba PDF." in combined
            assert all("page_num" in r for r in result)
        finally:
            Path(path).unlink(missing_ok=True)

    @pytest.mark.asyncio
    async def test_empty_pdf(self):
        """GIVEN an empty PDF WHEN parsed THEN result has empty text."""
        import fitz
        tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
        doc = fitz.open()
        doc.new_page()  # blank page
        doc.save(tmp.name)
        doc.close()
        try:
            result = await parse_pdf(tmp.name)
            assert len(result) >= 1
        finally:
            Path(tmp.name).unlink(missing_ok=True)


class TestParseDocx:
    """RED→GREEN→REFACTOR: DOCX text extraction."""

    @pytest.mark.asyncio
    async def test_extracts_text(self):
        """GIVEN a DOCX WHEN parsed THEN text is extracted."""
        path = _create_test_docx(["Primer párrafo.", "Segundo párrafo."])
        try:
            result = await parse_docx(path)
            assert len(result) >= 1
            combined = "\n".join(r["text"] for r in result)
            assert "Primer párrafo" in combined
            assert "Segundo párrafo" in combined
        finally:
            Path(path).unlink(missing_ok=True)

    @pytest.mark.asyncio
    async def test_returns_page_num_1(self):
        """DOCX has no native pages, so page_num is always 1."""
        path = _create_test_docx(["Solo un párrafo."])
        try:
            result = await parse_docx(path)
            assert result[0]["page_num"] == 1
        finally:
            Path(path).unlink(missing_ok=True)


class TestParseXlsx:
    """RED→GREEN→REFACTOR: XLSX text extraction."""

    @pytest.mark.asyncio
    async def test_extracts_sheet_text(self):
        """GIVEN an XLSX WHEN parsed THEN each sheet becomes a page."""
        sheets = {
            "Sheet1": [["Nombre", "Edad"], ["Ana", "30"]],
            "Sheet2": [["Producto", "Precio"], ["Zapatos", "15000"]],
        }
        path = _create_test_xlsx(sheets)
        try:
            result = await parse_xlsx(path)
            assert len(result) == 2
            sheet_names = {r["page_num"] for r in result}
            assert "Sheet1" in sheet_names
            assert "Sheet2" in sheet_names
        finally:
            Path(path).unlink(missing_ok=True)

    @pytest.mark.asyncio
    async def test_rows_joined_with_separator(self):
        """Rows in a sheet are joined with ' | '."""
        sheets = {
            "Data": [["A", "B"], ["C", "D"]],
        }
        path = _create_test_xlsx(sheets)
        try:
            result = await parse_xlsx(path)
            text = result[0]["text"]
            assert "A | B" in text
            assert "C | D" in text
        finally:
            Path(path).unlink(missing_ok=True)

    @pytest.mark.asyncio
    async def test_empty_sheet(self):
        """Empty sheet returns no text."""
        sheets = {"Empty": []}
        path = _create_test_xlsx(sheets)
        try:
            result = await parse_xlsx(path)
            assert result[0]["text"] == ""
        finally:
            Path(path).unlink(missing_ok=True)
